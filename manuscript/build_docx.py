#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""manuscript .md -> .docx (Times New Roman 12pt, heading levels, tables, figures).
遵循用户约定：cover/manuscript 生成 Times New Roman 12pt；不使用花哨模板。
v1.1.0 修订：修正图插入层级（### 3.x 也是 h3），支持 markdown 表格渲染（Table 1）。"""
import re, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

MD = "D:/2026.10/机械应力/manuscript/manuscript_v1.1.0.md"
FIG = "D:/2026.10/机械应力/analysis/results/layer1/fig"
OUT = "D:/2026.10/机械应力/manuscript/manuscript_v1.1.0.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)

def add_par(text, bold=False, size=12, align=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    r.bold = bold; r.italic = italic
    if align: p.alignment = align
    return p

def add_rich(text, size=12, bold_all=False):
    """处理 **bold** 与 *italic* 行内标记。"""
    p = doc.add_paragraph()
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*)", text)
    for t in tokens:
        if not t: continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            r = p.add_run(t)
        r.font.name = "Times New Roman"; r.font.size = Pt(size)
        if bold_all: r.bold = True
    return p

def heading_text(line):
    """提取标题纯文本（去 # 与首尾空白）。"""
    return line.lstrip("#").strip()

def is_table_row(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|")

def is_separator(line):
    s = line.strip()
    return bool(re.fullmatch(r"\|\s*-{2,}.*\|", s)) or set(s) <= set("|-: ")

def render_table(rows):
    """rows: list of raw '| a | b | c |' 字符串（不含分隔行）。"""
    # 以首行定列数
    cells = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    ncol = len(cells)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, c in enumerate(cells):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(c)
        run.bold = True; run.font.name = "Times New Roman"; run.font.size = Pt(10)
    for r in rows[1:]:
        vals = [c.strip() for c in r.strip().strip("|").split("|")]
        vals = (vals + [""] * ncol)[:ncol]
        cells_ = table.add_row().cells
        for i, v in enumerate(vals):
            cells_[i].text = ""
            run = cells_[i].paragraphs[0].add_run(v)
            run.font.name = "Times New Roman"; run.font.size = Pt(10)
    doc.add_paragraph()

def insert_figures(kind):
    if kind == "design":
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(os.path.join(FIG, "fig1_design.png"), width=Inches(6.2))
        add_par("Figure 1. Study design and prospectively specified analysis pipeline.", size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    elif kind == "dissociation":
        specs = [
            ("fig2a_module_scores.png", "Figure 2a. Module-score differences for the primary contrast (HZ-exo vs control-exo), with neutral controls displayed alongside. E1 is the only biological module significantly shifted."),
            ("fig2b_cci.png", "Figure 2b. Channel-complementarity index (CCI) differences by contrast (all non-significant)."),
            ("fig2c_E1_genes.png", "Figure 2c. Gene-level differences within E1 for the primary contrast: coherent suppression of classical nociceptor channels."),
            ("fig2d_dissociation.png", "Figure 2d. Perturbation route and channel programme: E1-selective suppression under HZ exosomes (interaction P = 0.0017); reciprocal VZV pattern not significant (P = 0.19)."),
        ]
        for fn, cap in specs:
            fp = os.path.join(FIG, fn)
            if os.path.exists(fp):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(fp, width=Inches(5.8))
                add_par(cap, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

lines = open(MD, encoding="utf-8").read().splitlines()
fig_state = {"design": False, "dissociation": False}
i = 0
first_line = True
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1; continue
    if is_table_row(line):
        # 收集连续表格行（跳过分隔行）
        tbl = []
        while i < len(lines) and is_table_row(lines[i]):
            if not is_separator(lines[i]):
                tbl.append(lines[i].rstrip())
            i += 1
        if len(tbl) >= 2:
            render_table(tbl)
        continue
    if line.startswith("# ") and first_line:
        add_par(heading_text(line), bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
        first_line = False
    elif line.startswith("# "):
        add_par(heading_text(line), bold=True, size=14)
    elif line.startswith("## "):
        h = heading_text(line)
        add_par(h, bold=True, size=12.5)
        if "3.1" in h and not fig_state["design"]:
            insert_figures("design"); fig_state["design"] = True
        if "3.4" in h and not fig_state["dissociation"]:
            insert_figures("dissociation"); fig_state["dissociation"] = True
    elif line.startswith("### "):
        h = heading_text(line)
        add_par(h, bold=True, size=12)
        if "3.1" in h and not fig_state["design"]:
            insert_figures("design"); fig_state["design"] = True
        if "3.4" in h and not fig_state["dissociation"]:
            insert_figures("dissociation"); fig_state["dissociation"] = True
    elif line.strip() == "---":
        pass
    elif line.startswith("> "):
        add_par(line[2:].strip(), size=10, italic=True)
    elif line.startswith("- "):
        add_rich("• " + line[2:].strip(), size=12)
    elif line.startswith("[") and re.match(r"^\[\d+\]", line):
        add_par(line.strip(), size=11)
    elif line.startswith("**") and line.rstrip().endswith("**") and len(line) < 80:
        add_par(line.strip("*"), bold=True, size=12)
    else:
        add_rich(line.strip(), size=12)
    i += 1

doc.save(OUT)
print("docx saved:", OUT)

# 字数统计（正文近似：摘要+正文，不含标题行/参考文献）
body = []
in_ref = False
for line in lines:
    s = line.strip()
    if s.startswith("## References"): in_ref = True
    if not s or s.startswith("#") or s == "---": continue
    if in_ref: continue
    body.append(s)
wc = len(" ".join(body).split())
print("main-text word count (incl. abstract, excl. refs):", wc)
